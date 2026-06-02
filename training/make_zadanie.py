"""
Генерация бланка «Задание на ВКР» (.docx) — точно по структуре бланка
кафедры ИСТ ЯрГУ (извлечён из присланного .doc), компактно на один лист.

Запуск:
    python training/make_zadanie.py

Артефакт:
    docs/zadanie-vkr.docx
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


STUDENT_FIO = "Журавлев Тарас Русланович"
TOPIC = ("Разработка системы автоматической классификации фотографий "
         "по визуальному стилю на основе transfer learning")
ADVISOR = "Д.Ю. Чалый"


def set_font(run, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:eastAsia', 'w:cs', 'w:ascii', 'w:hAnsi'):
        rfonts.set(qn(a), name)


def p(doc, text, size=12, bold=False, italic=False,
      align=WD_ALIGN_PARAGRAPH.LEFT, after=2, before=0):
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.line_spacing = 1.0
    if text:
        run = par.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic)
    return par


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(1.5)

    C = WD_ALIGN_PARAGRAPH.CENTER
    R = WD_ALIGN_PARAGRAPH.RIGHT
    J = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Шапка
    p(doc, "Федеральное государственное бюджетное образовательное учреждение",
      size=11, align=C)
    p(doc, "высшего образования", size=11, align=C)
    p(doc, "Ярославский государственный университет им. П.Г. Демидова",
      size=11, bold=True, align=C)
    p(doc, "Кафедра информационных и сетевых технологий", size=11, align=C, after=6)
    p(doc, "Направление подготовки 09.03.03 Прикладная информатика", size=11, align=C)
    p(doc, "Информационные технологии в цифровой экономике", size=11, align=C, after=10)

    # Утверждаю
    p(doc, "УТВЕРЖДАЮ", size=12, align=R)
    p(doc, "Зав. кафедрой, к.ф.-м.н.", size=12, align=R)
    p(doc, "_______________ Чалый Д.Ю.", size=12, align=R)
    p(doc, "«14» апреля 2026 г.", size=12, align=R, after=12)

    # Заголовок
    p(doc, "ЗАДАНИЕ", size=14, bold=True, align=C, after=2)
    p(doc, "по подготовке выпускной квалификационной работы студенту",
      size=12, align=C, after=2)
    p(doc, STUDENT_FIO, size=12, bold=True, align=C, after=12)

    # 1. Тема
    p(doc, "1. Тема работы", size=12, bold=True)
    p(doc, TOPIC, size=12, bold=True, align=J)
    p(doc, "Утверждена приказом № 354 от 13 апреля 2026 г.",
      size=11, italic=True, after=10)

    # 2. Срок сдачи
    p(doc, "2. Срок сдачи законченной работы: 1 июня 2026 г.",
      size=12, bold=True, after=10)

    # 3. Исходные данные
    p(doc, "3. Исходные данные к работе", size=12, bold=True)
    p(doc, "Датасет фотографий, собранный через открытые API стоковых сервисов "
           "Unsplash и Pexels; предобученная свёрточная нейронная сеть "
           "EfficientNet-B0; программные средства: Python (PyTorch, FastAPI, "
           "scikit-learn), Java (Spring Boot), React; инфраструктура: "
           "PostgreSQL, MinIO, RabbitMQ, Docker.",
      size=12, align=J, after=10)

    # 4. Перечень вопросов
    p(doc, "4. Перечень подлежащих разработке вопросов:", size=12, bold=True)
    for i, item in enumerate([
        "обзор методов классификации изображений, обоснование выбора "
        "архитектуры нейронной сети;",
        "сбор, разметка и предобработка датасета фотографий "
        "по стилевым категориям;",
        "обучение модели классификации методом transfer learning, "
        "оценка качества;",
        "проектирование и реализация микросервисной архитектуры "
        "(серверная часть, ML-сервис, веб-клиент);",
        "проведение исследовательских экспериментов, развёртывание "
        "и тестирование системы.",
    ], 1):
        p(doc, f"{i}) {item}", size=12, align=J, after=2)

    # Подписи
    doc.add_paragraph()
    p(doc, "Руководитель ВКР _______________ / Д.Ю. Чалый /", size=12, after=10)
    p(doc, "Задание принял к исполнению _______________ / Т.Р. Журавлев /",
      size=12, after=10)
    p(doc, "«14» апреля 2026 г.", size=12)

    out = Path(__file__).resolve().parent.parent / "docs" / "zadanie-vkr.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"saved: {out}")


if __name__ == "__main__":
    build()
